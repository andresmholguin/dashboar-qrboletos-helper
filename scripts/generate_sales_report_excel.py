import sys
import os
import json
import argparse
import datetime
import asyncio
from playwright.async_api import async_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def log(msg):
    sys.stderr.write(f"[INFO] {msg}\n")
    sys.stderr.flush()

def clean_currency(val_str):
    """Convierte un string como 'COP 82.305.000' o '$ 82.305.000' a número entero."""
    if not val_str:
        return 0
    tokens = str(val_str).replace("$", "").split("COP")
    target = tokens[-1].strip() if tokens else str(val_str)
    cleaned = target.replace(".", "").replace(",", "").strip()
    digits = "".join(ch for ch in cleaned if ch.isdigit() or ch == '-')
    if not digits:
        return 0
    try:
        return int(digits)
    except:
        return 0

def clean_int(val_str):
    if not val_str:
        return 0
    lines = [l.strip() for l in str(val_str).split("\n") if l.strip()]
    target = lines[-1] if lines else str(val_str)
    digits = "".join(ch for ch in target if ch.isdigit() or ch == '-')
    if not digits:
        return 0
    try:
        return int(digits)
    except:
        return 0

class SessionExpiredError(Exception):
    """Excepción lanzada cuando la sesión de QRBoletos en Chrome ha expirado."""
    pass

async def scrape_show_sales(page, show_url):
    """Navega a la URL de summary.aspx y extrae los datos de ventas estructurados."""
    if "/reports/sales/summary.aspx" not in show_url:
        if "/shows/" in show_url:
            parts = show_url.split("/shows/")[1].split("/")
            show_id = parts[0]
            base = show_url.split("/shows/")[0] + f"/shows/{show_id}"
            target_url = f"{base}/reports/sales/summary.aspx"
        else:
            target_url = show_url.rstrip("/") + "/reports/sales/summary.aspx"
    else:
        target_url = show_url

    log(f"Extrayendo informe de ventas: {target_url}")
    await page.goto(target_url, wait_until="domcontentloaded", timeout=20000)

    # Verificación instantánea de sesión activa en QRBoletos (con auto-login si ya tiene credenciales)
    curr_url = page.url.lower()
    if "login.aspx" in curr_url or "user/login" in curr_url:
        try:
            has_creds = await page.evaluate('''() => {
                const pass = document.querySelector("input[type='password']");
                return !!(pass && pass.value && pass.value.length > 0);
            }''')
            if has_creds:
                log("ℹ️ Credenciales recordadas detectadas en login.aspx. Haciendo clic automático en 'Iniciar sesión'...")
                login_btn = await page.query_selector("#login-button, button[type='submit'], input[type='submit'], .btn-primary")
                if login_btn:
                    await login_btn.click()
                    try:
                        await page.wait_for_load_state("domcontentloaded", timeout=12000)
                    except Exception:
                        pass
                    await asyncio.sleep(2)
                    if "login.aspx" not in page.url.lower():
                        log("✅ Auto-login exitoso. Redirigiendo a informe...")
                        await page.goto(target_url, wait_until="domcontentloaded", timeout=20000)
        except Exception as e_login:
            log(f"Aviso en intento de auto-login: {e_login}")

    curr_url = page.url.lower()
    if "login.aspx" in curr_url or "user/login" in curr_url:
        log("❌ [SESION EXPIRADA] Redirección a login.aspx detectada en Google Chrome.")
        raise SessionExpiredError("SESSION_EXPIRED: La sesión en QRBoletos ha caducado. Por favor inicia sesión en Google Chrome.")

    is_login_page = await page.evaluate('''() => {
        const hasUserInput = !!document.querySelector("#txtUsuario, input[type='password']");
        const hasLoginTitle = document.title && document.title.toLowerCase().includes("iniciar sesión");
        return hasUserInput || hasLoginTitle;
    }''')
    if is_login_page:
        log("❌ [SESION EXPIRADA] Formulario de inicio de sesión detectado en Google Chrome.")
        raise SessionExpiredError("SESSION_EXPIRED: La sesión en QRBoletos ha caducado. Por favor inicia sesión en Google Chrome.")

    await page.wait_for_selector(".table, table, h1, h2", timeout=9000)

    sales_data = await page.evaluate('''() => {
        // 1. Extraer metadatos del evento / espectáculo
        const meta = {};
        const metaRows = Array.from(document.querySelectorAll("table tr"));
        for (const r of metaRows) {
            const tds = Array.from(r.querySelectorAll("td"));
            if (tds.length >= 2) {
                const label = tds[0].innerText.trim();
                const val = tds[tds.length - 1].innerText.trim();
                if ((label === "Evento:" || (label.includes("Evento:") && !label.includes("Tipo") && !label.includes("PULEP"))) && !meta.evento) {
                    meta.evento = val;
                }
                if (label.includes("Espectáculo:") || label.includes("Espectaculo:")) {
                    meta.espectaculo = val;
                }
                if (label.includes("PULEP")) meta.pulep = val;
                if (label.includes("Sitio") || label.includes("Lugar")) meta.sitio = val;
                if (label.includes("Fecha de inicio")) meta.fechaInicio = val;
                if (label.includes("Fecha de finalización")) meta.fechaFin = val;
            }
        }

        // Si no se encontró evento, intentar con espectáculo o título de página
        if (!meta.evento && meta.espectaculo) {
            meta.evento = meta.espectaculo;
        }
        if (!meta.evento) {
            const titleEl = document.querySelector("h1, h2, h3, .page-title");
            if (titleEl) meta.evento = titleEl.innerText.trim();
        }

        // 2. Extraer tabla resumen por localidad
        let resumenLocalidades = [];

        const allTables = Array.from(document.querySelectorAll("table"));
        for (const t of allTables) {
            const ths = Array.from(t.querySelectorAll("th")).map(th => th.innerText.trim());
            const hasSummaryHeaders = ths.some(h => h.includes("Cantidad de entradas") || h.includes("Valor recaudado"));
            
            if (hasSummaryHeaders) {
                const trs = Array.from(t.querySelectorAll("tbody tr"));
                for (const tr of trs) {
                    const tds = Array.from(tr.querySelectorAll("td")).map(td => td.innerText.trim());
                    if (tds.length >= 4) {
                        const locName = tds[0];
                        const vendidas = tds[1];
                        const recaudo = tds[2];
                        const servicio = tds[3];

                        if (locName && vendidas && !locName.toLowerCase().includes("total")) {
                            resumenLocalidades.push({
                                localidad: locName,
                                vendidas: vendidas,
                                recaudoEntradas: recaudo,
                                recaudoServicio: servicio,
                                cortesias: 0,
                                boletosPagados: 0,
                                totalBoletos: 0
                            });
                        }
                    }
                }
            }
        }

        // Función para reconstruir matriz 2D con soporte exacto de rowspan y colspan
        function parseTableGrid(table) {
            const grid = [];
            const rows = Array.from(table.querySelectorAll('tbody tr'));
            for (let r = 0; r < rows.length; r++) {
                if (!grid[r]) grid[r] = [];
                const cells = Array.from(rows[r].querySelectorAll('td'));
                let cellIdx = 0;
                let c = 0;
                while (cellIdx < cells.length) {
                    while (grid[r][c] !== undefined) {
                        c++;
                    }
                    const cell = cells[cellIdx++];
                    const rowSpan = parseInt(cell.getAttribute('rowspan') || '1', 10);
                    const colSpan = parseInt(cell.getAttribute('colspan') || '1', 10);
                    let text = cell.innerText.trim();

                    const priceEl = cell.querySelector('.price');
                    if (priceEl) {
                        const oldPrice = priceEl.querySelector('.price-old');
                        if (oldPrice) {
                            const clone = priceEl.cloneNode(true);
                            const oldClone = clone.querySelector('.price-old');
                            if (oldClone) oldClone.remove();
                            text = clone.innerText.trim();
                        }
                    }

                    for (let i = 0; i < rowSpan; i++) {
                        if (!grid[r + i]) grid[r + i] = [];
                        for (let j = 0; j < colSpan; j++) {
                            grid[r + i][c + j] = text;
                        }
                    }
                    c += colSpan;
                }
            }
            return grid;
        }

        // 3. Extraer desglose detallado por localidad y canal de venta
        const channelTables = Array.from(document.querySelectorAll("table[id*='table-']"));
        const detalleCanales = [];

        for (const ct of channelTables) {
            const innerCard = ct.closest('.card');
            const outerCard = innerCard && innerCard.parentElement ? innerCard.parentElement.closest('.card') : null;
            const outerHeader = outerCard ? outerCard.querySelector('.card-header, h3, h4, h5') : null;
            let locHeader = outerHeader ? outerHeader.innerText.trim().split('\\n')[0].trim() : "";
            if (!locHeader) {
                locHeader = ct.id.replace(/^table-/, '').replace(/-cop$/, '').replace(/-/g, ' ').toUpperCase();
            }

            const grid = parseTableGrid(ct);
            for (const row of grid) {
                if (row.length >= 9) {
                    const etapa = row[0] || "";
                    const ref = row[1] || "";
                    const canal = row[2] || "";
                    const cupon = row[3] || "N/A";
                    const cantidad = row[4] || "0";
                    const valorEntrada = row[5] || "0";
                    const totalEntradas = row[6] || "0";
                    const valorServicio = row[7] || "0";
                    const totalServicio = row[8] || "0";

                    const isCortesia = ref.toUpperCase().includes("CORTES") || 
                                       etapa.toUpperCase().includes("CORTES") || 
                                       canal.toUpperCase().includes("CORTES");

                    detalleCanales.push({
                        localidad: locHeader,
                        etapa,
                        referencia: ref,
                        canal,
                        cupon,
                        cantidad,
                        valorEntrada,
                        totalEntradas,
                        valorServicio,
                        totalServicio,
                        isCortesia
                    });
                }
            }
        }

        // 4. Calcular cortesías y boletos pagados por localidad
        for (const loc of resumenLocalidades) {
            const locNameUpper = (loc.localidad || "").trim().toUpperCase();
            let cortesiasCount = 0;
            let pagadosCount = 0;

            for (const d of detalleCanales) {
                const dLocUpper = (d.localidad || "").trim().toUpperCase();
                if (dLocUpper === locNameUpper || dLocUpper.includes(locNameUpper) || locNameUpper.includes(dLocUpper)) {
                    const qty = parseInt(String(d.cantidad).replace(/[^0-9]/g, '') || '0', 10);
                    if (d.isCortesia) {
                        cortesiasCount += qty;
                    } else {
                        pagadosCount += qty;
                    }
                }
            }

            const totalVendidas = parseInt(String(loc.vendidas).replace(/[^0-9]/g, '') || '0', 10);
            loc.cortesias = cortesiasCount;
            loc.boletosPagados = Math.max(0, totalVendidas - cortesiasCount);
            loc.totalBoletos = totalVendidas;
        }

        return {
            meta,
            resumenLocalidades,
            detalleCanales,
            url: window.location.href
        };
    }''')

    # 5. Extraer datos de 'Recaudo en ventas' (collected.aspx)
    collected_url = target_url.replace("/summary.aspx", "/collected.aspx")
    log(f"Consultando recaudo en ventas: {collected_url}")
    recaudo_info = {"metodos": [], "entregaEmpresario": None}
    try:
        await page.goto(collected_url, wait_until="domcontentloaded", timeout=15000)
        recaudo_info = await page.evaluate('''() => {
            const result = {
                metodos: [],
                entregaEmpresario: null
            };
            const tables = Array.from(document.querySelectorAll("table"));
            for (const t of tables) {
                const text = t.innerText || "";
                if (text.includes("Método de pago") || text.includes("Metodo de pago") || text.includes("Mtodo de pago")) {
                    const rows = Array.from(t.querySelectorAll("tr"));
                    for (const r of rows) {
                        const cells = Array.from(r.querySelectorAll("th, td")).map(c => c.innerText.trim());
                        if (cells.length >= 2) {
                            const metodo = cells[0];
                            const totalStr = cells[cells.length - 1];
                            if (metodo && !metodo.toLowerCase().includes("método") && !metodo.toLowerCase().includes("metodo") && !metodo.toLowerCase().includes("mtodo")) {
                                result.metodos.push({ metodo, totalStr });
                                const cleanMetodo = metodo.toLowerCase().normalize("NFD").replace(/[\u0300-\u036f]/g, "");
                                if (cleanMetodo.includes("entrega a empresario")) {
                                    result.entregaEmpresario = {
                                        tiene: true,
                                        metodo: metodo,
                                        totalStr: totalStr
                                    };
                                }
                            }
                        }
                    }
                }
            }
            return result;
        }''')
    except Exception as e:
        log(f"Aviso al consultar collected.aspx: {e}")

    if recaudo_info.get("entregaEmpresario"):
        raw_val = recaudo_info["entregaEmpresario"].get("totalStr", "")
        clean_val = clean_currency(raw_val)
        recaudo_info["entregaEmpresario"]["total"] = clean_val
        recaudo_info["entregaEmpresario"]["totalFormatted"] = f"COP {clean_val:,}".replace(",", ".")
        log(f"✅ 'Entrega a empresario' detectada: {recaudo_info['entregaEmpresario']['totalFormatted']}")

    sales_data["recaudoMetodos"] = recaudo_info.get("metodos", [])
    sales_data["entregaEmpresario"] = recaudo_info.get("entregaEmpresario")

    return sales_data

def build_excel_report(all_events_sales, output_path):
    """Construye el libro Excel con estilos ejecutivos usando openpyxl."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Estilos corporativos
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    accent_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    cortesia_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    total_border = Border(
        top=Side(style="thin", color="0F172A"),
        bottom=Side(style="double", color="0F172A")
    )

    currency_format = '"$"#,##0'
    number_format = '#,##0'

    # ----------------------------------------------------
    # HOJA 1: RESUMEN GENERAL DE EVENTOS
    # ----------------------------------------------------
    ws_general = wb.create_sheet(title="Resumen General")
    ws_general.views.sheetView[0].showGridLines = True

    ws_general["A1"] = "REPORTE CONSOLIDADO DE VENTAS Y CORTESÍAS - QRBOLETOS"
    ws_general["A1"].font = title_font
    ws_general["A2"] = f"Generado el: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total eventos analizados: {len(all_events_sales)}"
    ws_general["A2"].font = subtitle_font

    headers_general = [
        "Evento",
        "Espectáculo / Fecha",
        "Fecha",
        "Sitio / Lugar",
        "PULEP",
        "Boletos Pagados",
        "Cortesías Emitidas",
        "Total Boletos",
        "Recaudo Entradas (COP)",
        "Cover Service (COP)",
        "Total Recaudado (COP)",
        "Entrega Empresario (COP)"
    ]

    start_row = 4
    for col_num, h in enumerate(headers_general, 1):
        cell = ws_general.cell(row=start_row, column=col_num, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num > 5 else "left", vertical="center")

    curr_row = start_row + 1
    for ev in all_events_sales:
        meta = ev.get("meta", {})
        resumen = ev.get("resumenLocalidades", [])

        ev_pagados = sum(l.get("boletosPagados", 0) for l in resumen)
        ev_cortesias = sum(l.get("cortesias", 0) for l in resumen)
        ev_total_b = sum(clean_int(l.get("totalBoletos", l.get("vendidas"))) for l in resumen)
        ev_recaudo = sum(clean_currency(l.get("recaudoEntradas")) for l in resumen)
        ev_servicio = sum(clean_currency(l.get("recaudoServicio")) for l in resumen)
        ev_total = ev_recaudo + ev_servicio

        ev_ee = ev.get("entregaEmpresario")
        ee_val = ev_ee.get("total", 0) if ev_ee else 0

        ws_general.cell(row=curr_row, column=1, value=meta.get("evento", "Evento")).font = bold_font
        ws_general.cell(row=curr_row, column=2, value=meta.get("espectaculo", "")).font = normal_font
        ws_general.cell(row=curr_row, column=3, value=meta.get("fechaInicio", "")).font = normal_font
        ws_general.cell(row=curr_row, column=4, value=meta.get("sitio", "")).font = normal_font
        ws_general.cell(row=curr_row, column=5, value=meta.get("pulep", "")).font = normal_font

        c6 = ws_general.cell(row=curr_row, column=6, value=ev_pagados)
        c6.number_format = number_format
        c6.font = normal_font
        c6.alignment = Alignment(horizontal="right")

        c7 = ws_general.cell(row=curr_row, column=7, value=ev_cortesias)
        c7.number_format = number_format
        c7.font = bold_font if ev_cortesias > 0 else normal_font
        if ev_cortesias > 0:
            c7.fill = cortesia_fill
        c7.alignment = Alignment(horizontal="right")

        c8 = ws_general.cell(row=curr_row, column=8, value=ev_total_b)
        c8.number_format = number_format
        c8.font = bold_font
        c8.alignment = Alignment(horizontal="right")

        c9 = ws_general.cell(row=curr_row, column=9, value=ev_recaudo)
        c9.number_format = currency_format
        c9.font = normal_font
        c9.alignment = Alignment(horizontal="right")

        c10 = ws_general.cell(row=curr_row, column=10, value=ev_servicio)
        c10.number_format = currency_format
        c10.font = normal_font
        c10.alignment = Alignment(horizontal="right")

        c11 = ws_general.cell(row=curr_row, column=11, value=ev_total)
        c11.number_format = currency_format
        c11.font = bold_font
        c11.alignment = Alignment(horizontal="right")

        c12 = ws_general.cell(row=curr_row, column=12, value=ee_val)
        c12.number_format = currency_format
        c12.font = bold_font if ee_val > 0 else normal_font
        if ee_val > 0:
            c12.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        c12.alignment = Alignment(horizontal="right")

        for c in range(1, 13):
            ws_general.cell(row=curr_row, column=c).border = thin_border

        curr_row += 1

    # Fila de Totales Generales
    if curr_row > start_row + 1:
        total_row = curr_row
        ws_general.cell(row=total_row, column=1, value="TOTAL CONSOLIDADO").font = bold_font
        ws_general.cell(row=total_row, column=6, value=f"=SUM(F{start_row+1}:F{total_row-1})").number_format = number_format
        ws_general.cell(row=total_row, column=7, value=f"=SUM(G{start_row+1}:G{total_row-1})").number_format = number_format
        ws_general.cell(row=total_row, column=8, value=f"=SUM(H{start_row+1}:H{total_row-1})").number_format = number_format
        ws_general.cell(row=total_row, column=9, value=f"=SUM(I{start_row+1}:I{total_row-1})").number_format = currency_format
        ws_general.cell(row=total_row, column=10, value=f"=SUM(J{start_row+1}:J{total_row-1})").number_format = currency_format
        ws_general.cell(row=total_row, column=11, value=f"=SUM(K{start_row+1}:K{total_row-1})").number_format = currency_format
        ws_general.cell(row=total_row, column=12, value=f"=SUM(L{start_row+1}:L{total_row-1})").number_format = currency_format

        for c in range(1, 13):
            cell = ws_general.cell(row=total_row, column=c)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = total_border
            if c >= 6:
                cell.alignment = Alignment(horizontal="right")

    # ----------------------------------------------------
    # HOJA 2: DETALLE POR LOCALIDAD Y CORTESÍAS
    # ----------------------------------------------------
    ws_loc = wb.create_sheet(title="Ventas por Localidad")
    ws_loc.views.sheetView[0].showGridLines = True

    ws_loc["A1"] = "DETALLE DE VENTAS Y CORTESÍAS POR LOCALIDAD"
    ws_loc["A1"].font = title_font

    headers_loc = [
        "Evento",
        "Espectáculo / Fecha",
        "Localidad",
        "Boletos Pagados",
        "Cortesías Emitidas",
        "Total Boletos",
        "Recaudo Entradas (COP)",
        "Cover Service (COP)",
        "Total Recaudo (COP)"
    ]

    for col_num, h in enumerate(headers_loc, 1):
        cell = ws_loc.cell(row=3, column=col_num, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num > 3 else "left", vertical="center")

    loc_row = 4
    for ev in all_events_sales:
        meta = ev.get("meta", {})
        ev_name = meta.get("evento", "Evento")
        esp_name = meta.get("espectaculo", "")
        resumen = ev.get("resumenLocalidades", [])

        for loc_item in resumen:
            l_pag = loc_item.get("boletosPagados", 0)
            l_cor = loc_item.get("cortesias", 0)
            l_tot_b = loc_item.get("totalBoletos", clean_int(loc_item.get("vendidas", 0)))
            v_rec = clean_currency(loc_item.get("recaudoEntradas"))
            v_ser = clean_currency(loc_item.get("recaudoServicio"))
            v_tot = v_rec + v_ser

            ws_loc.cell(row=loc_row, column=1, value=ev_name).font = normal_font
            ws_loc.cell(row=loc_row, column=2, value=esp_name).font = normal_font
            ws_loc.cell(row=loc_row, column=3, value=loc_item.get("localidad", "")).font = bold_font

            c4 = ws_loc.cell(row=loc_row, column=4, value=l_pag)
            c4.number_format = number_format
            c4.font = normal_font
            c4.alignment = Alignment(horizontal="right")

            c5 = ws_loc.cell(row=loc_row, column=5, value=l_cor)
            c5.number_format = number_format
            c5.font = bold_font if l_cor > 0 else normal_font
            if l_cor > 0:
                c5.fill = cortesia_fill
            c5.alignment = Alignment(horizontal="right")

            c6 = ws_loc.cell(row=loc_row, column=6, value=l_tot_b)
            c6.number_format = number_format
            c6.font = bold_font
            c6.alignment = Alignment(horizontal="right")

            c7 = ws_loc.cell(row=loc_row, column=7, value=v_rec)
            c7.number_format = currency_format
            c7.font = normal_font
            c7.alignment = Alignment(horizontal="right")

            c8 = ws_loc.cell(row=loc_row, column=8, value=v_ser)
            c8.number_format = currency_format
            c8.font = normal_font
            c8.alignment = Alignment(horizontal="right")

            c9 = ws_loc.cell(row=loc_row, column=9, value=v_tot)
            c9.number_format = currency_format
            c9.font = bold_font
            c9.alignment = Alignment(horizontal="right")

            for col in range(1, 10):
                ws_loc.cell(row=loc_row, column=col).border = thin_border

            loc_row += 1

    if loc_row > 4:
        ws_loc.cell(row=loc_row, column=1, value="TOTAL GENERAL").font = bold_font
        ws_loc.cell(row=loc_row, column=4, value=f"=SUM(D4:D{loc_row-1})").number_format = number_format
        ws_loc.cell(row=loc_row, column=5, value=f"=SUM(E4:E{loc_row-1})").number_format = number_format
        ws_loc.cell(row=loc_row, column=6, value=f"=SUM(F4:F{loc_row-1})").number_format = number_format
        ws_loc.cell(row=loc_row, column=7, value=f"=SUM(G4:G{loc_row-1})").number_format = currency_format
        ws_loc.cell(row=loc_row, column=8, value=f"=SUM(H4:H{loc_row-1})").number_format = currency_format
        ws_loc.cell(row=loc_row, column=9, value=f"=SUM(I4:I{loc_row-1})").number_format = currency_format
        for col in range(1, 10):
            cell = ws_loc.cell(row=loc_row, column=col)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = total_border
            if col >= 4:
                cell.alignment = Alignment(horizontal="right")

    # ----------------------------------------------------
    # HOJA 3: DESGLOSE POR CANAL (WEB VS POS)
    # ----------------------------------------------------
    ws_chan = wb.create_sheet(title="Canales de Venta")
    ws_chan.views.sheetView[0].showGridLines = True

    ws_chan["A1"] = "DESGLOSE POR CANAL DE VENTA (SITIO WEB VS PUNTO DE VENTA)"
    ws_chan["A1"].font = title_font

    headers_chan = [
        "Evento",
        "Espectáculo / Fecha",
        "Localidad",
        "Etapa",
        "Referencia",
        "Canal de Venta",
        "Tipo Entrada",
        "Cantidad Vendida",
        "Valor Entrada",
        "Total Entradas (COP)",
        "Valor Servicio",
        "Total Servicio (COP)"
    ]

    for col_num, h in enumerate(headers_chan, 1):
        cell = ws_chan.cell(row=3, column=col_num, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num >= 8 else "left", vertical="center")

    chan_row = 4
    for ev in all_events_sales:
        meta = ev.get("meta", {})
        ev_name = meta.get("evento", "Evento")
        esp_name = meta.get("espectaculo", "")
        detalles = ev.get("detalleCanales", [])

        for d in detalles:
            c_cant = clean_int(d.get("cantidad"))
            c_v_ent = clean_currency(d.get("valorEntrada"))
            c_t_ent = clean_currency(d.get("totalEntradas"))
            c_v_ser = clean_currency(d.get("valorServicio"))
            c_t_ser = clean_currency(d.get("totalServicio"))
            is_cor = d.get("isCortesia", False)

            ws_chan.cell(row=chan_row, column=1, value=ev_name).font = normal_font
            ws_chan.cell(row=chan_row, column=2, value=esp_name).font = normal_font
            ws_chan.cell(row=chan_row, column=3, value=d.get("localidad", "")).font = bold_font
            ws_chan.cell(row=chan_row, column=4, value=d.get("etapa", "")).font = normal_font
            ws_chan.cell(row=chan_row, column=5, value=d.get("referencia", "")).font = normal_font
            
            c_canal = ws_chan.cell(row=chan_row, column=6, value=d.get("canal", ""))
            c_canal.font = bold_font
            if "web" in d.get("canal", "").lower():
                c_canal.fill = accent_fill

            c_tipo = ws_chan.cell(row=chan_row, column=7, value="CORTESÍA" if is_cor else "PAGADO")
            c_tipo.font = bold_font
            if is_cor:
                c_tipo.fill = cortesia_fill

            ws_chan.cell(row=chan_row, column=8, value=c_cant).number_format = number_format
            ws_chan.cell(row=chan_row, column=9, value=c_v_ent).number_format = currency_format
            ws_chan.cell(row=chan_row, column=10, value=c_t_ent).number_format = currency_format
            ws_chan.cell(row=chan_row, column=11, value=c_v_ser).number_format = currency_format
            ws_chan.cell(row=chan_row, column=12, value=c_t_ser).number_format = currency_format

            for col in range(1, 13):
                cell = ws_chan.cell(row=chan_row, column=col)
                cell.border = thin_border
                if col >= 8:
                    cell.alignment = Alignment(horizontal="right")

            chan_row += 1

    if chan_row > 4:
        ws_chan.cell(row=chan_row, column=1, value="TOTAL GENERAL").font = bold_font
        ws_chan.cell(row=chan_row, column=8, value=f"=SUM(H4:H{chan_row-1})").number_format = number_format
        ws_chan.cell(row=chan_row, column=10, value=f"=SUM(J4:J{chan_row-1})").number_format = currency_format
        ws_chan.cell(row=chan_row, column=12, value=f"=SUM(L4:L{chan_row-1})").number_format = currency_format
        for col in range(1, 13):
            cell = ws_chan.cell(row=chan_row, column=col)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = total_border
            if col >= 8:
                cell.alignment = Alignment(horizontal="right")

    # ----------------------------------------------------
    # HOJA 4: MÉTODOS DE PAGO Y RECAUDO EN VENTAS
    # ----------------------------------------------------
    ws_pay = wb.create_sheet(title="Métodos de Pago")
    ws_pay.views.sheetView[0].showGridLines = True

    ws_pay["A1"] = "DETALLE DE MÉTODOS DE PAGO Y RECAUDO EN VENTAS (COLLECTED)"
    ws_pay["A1"].font = title_font

    headers_pay = [
        "Evento",
        "Espectáculo",
        "Método de Pago",
        "Total Recaudado (COP)",
        "¿Entrega a Empresario?"
    ]

    for col_num, h in enumerate(headers_pay, 1):
        cell = ws_pay.cell(row=3, column=col_num, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num >= 4 else "left", vertical="center")

    pay_row = 4
    for ev in all_events_sales:
        meta = ev.get("meta", {})
        ev_name = meta.get("evento", "Evento")
        esp_name = meta.get("espectaculo", "")
        metodos = ev.get("recaudoMetodos", [])

        for m in metodos:
            m_nombre = m.get("metodo", "")
            m_val = clean_currency(m.get("totalStr", "0"))
            is_ee = "entrega a empresario" in m_nombre.lower()

            ws_pay.cell(row=pay_row, column=1, value=ev_name).font = normal_font
            ws_pay.cell(row=pay_row, column=2, value=esp_name).font = normal_font
            ws_pay.cell(row=pay_row, column=3, value=m_nombre).font = bold_font if is_ee else normal_font

            c_val = ws_pay.cell(row=pay_row, column=4, value=m_val)
            c_val.number_format = currency_format
            c_val.font = bold_font if is_ee else normal_font
            c_val.alignment = Alignment(horizontal="right")

            c_flag = ws_pay.cell(row=pay_row, column=5, value="SÍ (ENTREGA A EMPRESARIO)" if is_ee else "No")
            c_flag.font = bold_font if is_ee else normal_font
            c_flag.alignment = Alignment(horizontal="center")

            if is_ee:
                c_val.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                c_flag.fill = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")

            for c in range(1, 6):
                ws_pay.cell(row=pay_row, column=c).border = thin_border

            pay_row += 1

    # Autoajustar ancho de columnas para todas las hojas
    for ws in [ws_general, ws_loc, ws_chan, ws_pay]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and "\n" not in val_str:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)
    log(f"Libro Excel generado exitosamente en: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Generar informe de ventas de QRBoletos en Excel")
    parser.add_argument("--events-json", help="Ruta al archivo JSON con la lista de eventos o array JSON")
    parser.add_argument("--target-url", help="URL de un evento individual para informe único")
    parser.add_argument("--output-excel", required=True, help="Ruta de destino del archivo .xlsx")
    parser.add_argument("--cdp-url", default="http://localhost:9222", help="URL de Chrome CDP")
    parser.add_argument("--json-out", action="store_true", help="Imprimir datos extraídos en JSON a stdout")
    args = parser.parse_args()

    # Cargar lista de eventos a procesar
    events_to_process = []
    if args.target_url:
        events_to_process.append({"urlBase": args.target_url})
    elif args.events_json:
        try:
            if os.path.exists(args.events_json):
                with open(args.events_json, "r", encoding="utf-8") as f:
                    events_to_process = json.load(f)
            else:
                events_to_process = json.loads(args.events_json)
        except Exception as e:
            log(f"Error cargando events-json: {e}")
            sys.exit(1)

    if not events_to_process:
        log("No se proporcionaron eventos para procesar.")
        sys.exit(1)

    asyncio.run(async_main(args, events_to_process))

async def async_main(args, events_to_process):
    results = []
    session_expired_flag = False

    async with async_playwright() as p:
        try:
            browser = await p.chromium.connect_over_cdp(args.cdp_url)
        except Exception as e:
            log(f"No se pudo conectar a Chrome en {args.cdp_url}: {e}")
            if args.json_out:
                print(json.dumps({
                    "success": False,
                    "code": "CHROME_OFFLINE",
                    "error": f"Chrome no responde en {args.cdp_url}. Verifica que esté abierto con --remote-debugging-port=9222."
                }))
            sys.exit(1)

        context = browser.contexts[0] if browser.contexts else browser.new_context()

        # Chequeo preventivo: verificar si alguna pestaña abierta de QRBoletos está en login y si tiene credenciales
        for existing_page in context.pages:
            try:
                ep_url = existing_page.url.lower()
                if "qrboletos.com" in ep_url and ("login.aspx" in ep_url or "user/login" in ep_url):
                    has_creds = await existing_page.evaluate('''() => {
                        const pass = document.querySelector("input[type='password']");
                        return !!(pass && pass.value && pass.value.length > 0);
                    }''')
                    if has_creds:
                        log("ℹ️ Credenciales recordadas detectadas en pestaña de login. Haciendo clic automático en 'Iniciar sesión'...")
                        login_btn = await existing_page.query_selector("#login-button, button[type='submit'], input[type='submit'], .btn-primary")
                        if login_btn:
                            await login_btn.click()
                            try:
                                await existing_page.wait_for_load_state("domcontentloaded", timeout=12000)
                            except Exception:
                                pass
                            await asyncio.sleep(2)

                    if "login.aspx" in existing_page.url.lower() or "user/login" in existing_page.url.lower():
                        log("❌ [SESION EXPIRADA] Se detectó una pestaña de QRBoletos en la pantalla de inicio de sesión.")
                        if args.json_out:
                            print(json.dumps({
                                "success": False,
                                "code": "SESSION_EXPIRED",
                                "error": "Tu sesión en Google Chrome ha expirado. Inicia sesión en dashboard.qrboletos.com y vuelve a intentarlo."
                            }))
                        sys.exit(41)
            except Exception:
                pass

        queue = asyncio.Queue()
        for idx, ev in enumerate(events_to_process):
            await queue.put((idx, ev))

        concurrency = min(3, max(1, len(events_to_process)))

        async def worker():
            nonlocal session_expired_flag
            page = await context.new_page()
            try:
                while True:
                    if session_expired_flag:
                        break
                    try:
                        idx, ev = queue.get_nowait()
                    except asyncio.QueueEmpty:
                        break
                    url = ev.get("urlBase") or ev.get("url")
                    if not url:
                        queue.task_done()
                        continue
                    try:
                        data = await scrape_show_sales(page, url)
                        results.append((idx, data))
                    except SessionExpiredError as se:
                        session_expired_flag = True
                        log(f"❌ Sesión expirada durante extracción de {url}: {se}")
                        # Vaciar la cola para detener a los demás workers
                        while not queue.empty():
                            try:
                                queue.get_nowait()
                                queue.task_done()
                            except asyncio.QueueEmpty:
                                break
                        break
                    except Exception as e:
                        log(f"Error extrayendo {url}: {e}")
                    finally:
                        queue.task_done()
            finally:
                await page.close()

        workers = [asyncio.create_task(worker()) for _ in range(concurrency)]
        await queue.join()
        for w in workers:
            w.cancel()

    if session_expired_flag:
        log("❌ Operación abortada: La sesión en Google Chrome ha expirado.")
        if args.json_out:
            print(json.dumps({
                "success": False,
                "code": "SESSION_EXPIRED",
                "error": "Tu sesión en Google Chrome ha expirado. Inicia sesión en dashboard.qrboletos.com y vuelve a intentarlo."
            }))
        sys.exit(41)

    if not results:
        log("No se pudo obtener datos de ventas de ningún evento.")
        sys.exit(1)

    results.sort(key=lambda x: x[0])
    all_sales_data = [r[1] for r in results]

    # Generar Excel
    build_excel_report(all_sales_data, args.output_excel)

    if args.json_out:
        print(json.dumps({
            "success": True,
            "totalEvents": len(all_sales_data),
            "excelPath": args.output_excel,
            "salesData": all_sales_data
        }, ensure_ascii=False))

if __name__ == "__main__":
    main()
